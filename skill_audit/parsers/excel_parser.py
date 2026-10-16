"""Parser for Excel spreadsheets (.xlsx, .xlsm) and CSV files."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from skill_audit.grounding import extract_key_entities
from skill_audit.models import AuditDocument, DocType, KnowledgeChunk
from skill_audit.parsers.base import BaseParser


class ExcelParser(BaseParser):
    """Parses Excel spreadsheets and CSV tables into AuditDocument with row/col locators."""

    SUPPORTED_EXTS = {".xlsx", ".xlsm", ".xltx", ".csv"}

    def can_parse(self, path: Path) -> bool:
        return path.suffix.lower() in self.SUPPORTED_EXTS

    def parse(self, path: Path) -> AuditDocument:
        ext = path.suffix.lower()
        if ext == ".csv":
            return self._parse_csv(path)
        else:
            return self._parse_excel(path)

    def _parse_csv(self, path: Path) -> AuditDocument:
        rows: list[list[str]] = []
        with open(path, mode="r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            for r in reader:
                rows.append([str(c).strip() for c in r])

        rendered_lines: list[str] = []
        chunks: list[KnowledgeChunk] = []
        header = rows[0] if rows else []

        batch_rows: list[str] = []
        start_row = 1
        chunk_idx = 0

        for r_idx, row in enumerate(rows, 1):
            if not any(row):
                continue
            if header and r_idx > 1 and len(row) == len(header):
                row_str = f"Baris {r_idx}: " + " | ".join(f"{h}: {val}" for h, val in zip(header, row) if val)
            else:
                row_str = f"Baris {r_idx}: " + " | ".join(row)

            rendered_lines.append(row_str)
            batch_rows.append(row_str)

            if len(batch_rows) >= 25:
                chunk_text = "\n".join(batch_rows)
                locator = f"Baris {start_row}-{r_idx}"
                chunks.append(
                    KnowledgeChunk(
                        content=chunk_text,
                        document_name=path.name,
                        locator=locator,
                        chunk_index=chunk_idx,
                        metadata={"start_row": start_row, "end_row": r_idx},
                    )
                )
                chunk_idx += 1
                batch_rows = []
                start_row = r_idx + 1

        if batch_rows:
            chunk_text = "\n".join(batch_rows)
            locator = f"Baris {start_row}-{len(rows)}"
            chunks.append(
                KnowledgeChunk(
                    content=chunk_text,
                    document_name=path.name,
                    locator=locator,
                    chunk_index=chunk_idx,
                    metadata={"start_row": start_row, "end_row": len(rows)},
                )
            )

        full_content = "\n".join(rendered_lines)
        entities = extract_key_entities(full_content)

        return AuditDocument(
            name=path.name,
            path=str(path.resolve()),
            doc_type=DocType.CSV,
            total_pages_or_sheets=1,
            content=full_content,
            entities=entities,
            chunks=chunks,
        )

    def _parse_excel(self, path: Path) -> AuditDocument:
        import openpyxl

        wb = openpyxl.load_workbook(str(path), data_only=True)
        sheet_names = wb.sheetnames

        rendered_lines: list[str] = []
        chunks: list[KnowledgeChunk] = []
        chunk_idx = 0

        for sheet_name in sheet_names:
            ws = wb[sheet_name]
            rows_data = list(ws.iter_rows(values_only=True))
            if not rows_data:
                continue

            rendered_lines.append(f"=== [Lembar Kerja / Sheet: {sheet_name}] ===")
            header = [str(c).strip() if c is not None else f"Kolom_{i+1}" for i, c in enumerate(rows_data[0])]

            sheet_batch: list[str] = []
            start_row = 1

            for r_idx, row in enumerate(rows_data, 1):
                clean_row = [str(c).strip() if c is not None else "" for c in row]
                if not any(clean_row):
                    continue

                if r_idx > 1 and len(clean_row) <= len(header):
                    pairs = [f"{header[i]}: {val}" for i, val in enumerate(clean_row) if val]
                    row_str = f"[{sheet_name} : Baris {r_idx}] " + " | ".join(pairs)
                else:
                    row_str = f"[{sheet_name} : Baris {r_idx}] " + " | ".join(filter(None, clean_row))

                rendered_lines.append(row_str)
                sheet_batch.append(row_str)

                if len(sheet_batch) >= 20:
                    chunk_text = "\n".join(sheet_batch)
                    locator = f"{sheet_name} : Baris {start_row}-{r_idx}"
                    chunks.append(
                        KnowledgeChunk(
                            content=chunk_text,
                            document_name=path.name,
                            locator=locator,
                            chunk_index=chunk_idx,
                            metadata={"sheet": sheet_name, "start_row": start_row, "end_row": r_idx},
                        )
                    )
                    chunk_idx += 1
                    sheet_batch = []
                    start_row = r_idx + 1

            if sheet_batch:
                chunk_text = "\n".join(sheet_batch)
                locator = f"{sheet_name} : Baris {start_row}-{len(rows_data)}"
                chunks.append(
                    KnowledgeChunk(
                        content=chunk_text,
                        document_name=path.name,
                        locator=locator,
                        chunk_index=chunk_idx,
                        metadata={"sheet": sheet_name, "start_row": start_row, "end_row": len(rows_data)},
                    )
                )
                chunk_idx += 1

        full_content = "\n".join(rendered_lines)
        entities = extract_key_entities(full_content)

        return AuditDocument(
            name=path.name,
            path=str(path.resolve()),
            doc_type=DocType.EXCEL,
            total_pages_or_sheets=len(sheet_names),
            content=full_content,
            entities=entities,
            chunks=chunks,
        )
