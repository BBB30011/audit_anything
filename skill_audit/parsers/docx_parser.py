"""Parser for Word documents (.docx) using python-docx."""

from __future__ import annotations

from pathlib import Path

from skill_audit.grounding import extract_key_entities
from skill_audit.models import AuditDocument, DocType, KnowledgeChunk
from skill_audit.parsers.base import BaseParser


class DocxParser(BaseParser):
    """Parses .docx documents into AuditDocument with paragraph and table locators."""

    def can_parse(self, path: Path) -> bool:
        return path.suffix.lower() == ".docx"

    def parse(self, path: Path) -> AuditDocument:
        import docx

        doc = docx.Document(str(path))
        rendered_lines: list[str] = []
        chunks: list[KnowledgeChunk] = []

        chunk_idx = 0
        current_section = "Pendahuluan"

        for p_idx, p in enumerate(doc.paragraphs, 1):
            text = p.text.strip()
            if not text:
                continue

            # Detect heading or clause structure
            if p.style.name.startswith("Heading") or text.lower().startswith(("pasal ", "bab ", "article ", "section ")):
                current_section = text

            line_str = f"[Paragraf {p_idx}] {text}"
            rendered_lines.append(line_str)

        # Also extract tables
        for t_idx, table in enumerate(doc.tables, 1):
            table_rows: list[str] = []
            for r_idx, row in enumerate(table.rows, 1):
                cell_vals = [c.text.strip() for c in row.cells]
                if any(cell_vals):
                    table_rows.append(f"Baris {r_idx}: " + " | ".join(cell_vals))

            if table_rows:
                rendered_lines.append(f"\n--- [Tabel {t_idx}] ---\n" + "\n".join(table_rows) + "\n")
                chunks.append(
                    KnowledgeChunk(
                        content="\n".join(table_rows),
                        document_name=path.name,
                        section=f"Tabel {t_idx}",
                        locator=f"Tabel {t_idx}",
                        chunk_index=chunk_idx,
                        metadata={"table_index": t_idx},
                    )
                )
                chunk_idx += 1

        full_content = "\n".join(rendered_lines)
        entities = extract_key_entities(full_content)

        # Batch text into chunks
        paragraphs = [p for p in rendered_lines if not p.startswith("---")]
        batch: list[str] = []
        start_p = 1
        for i, line in enumerate(paragraphs, 1):
            batch.append(line)
            if len(batch) >= 15:
                chunk_text = "\n".join(batch)
                chunks.append(
                    KnowledgeChunk(
                        content=chunk_text,
                        document_name=path.name,
                        locator=f"Paragraf {start_p}-{i}",
                        chunk_index=chunk_idx,
                        metadata={"start_para": start_p, "end_para": i},
                    )
                )
                chunk_idx += 1
                batch = []
                start_p = i + 1

        if batch:
            chunk_text = "\n".join(batch)
            chunks.append(
                KnowledgeChunk(
                    content=chunk_text,
                    document_name=path.name,
                    locator=f"Paragraf {start_p}-{len(paragraphs)}",
                    chunk_index=chunk_idx,
                    metadata={"start_para": start_p, "end_para": len(paragraphs)},
                )
            )

        return AuditDocument(
            name=path.name,
            path=str(path.resolve()),
            doc_type=DocType.DOCX,
            total_pages_or_sheets=max(1, len(rendered_lines) // 30),
            content=full_content,
            entities=entities,
            chunks=chunks,
        )
